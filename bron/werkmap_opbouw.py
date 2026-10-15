"""Bouwt de werkmap Opioïdrotatie (basis-xlsx; VBA en knoppen komen erbij via Bouw_werkmap.ps1).

Opmaak in BENU-huisstijl: koningsblauw #002664, groen #58A618 / #407C05, lichtgrijsblauw #EEF2F7.
Lettertype Segoe UI (Ubuntu, het BENU-webfont, staat niet standaard op Windows).

LET OP: Pharmacom.bas gebruikt de cellen B4 (patiëntregel) en B/C/D/I van de rijen
R0..R1 (blok 1). Verplaats je die, pas dan R_EERST/R_LAATST in Pharmacom.bas aan.
"""
import sys
from copy import copy

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else "Opioidrotatie_basis.xlsx"
VERSIE = "Versie 1.7 - 15-10-2026"

# ------------------------------------------------------------------ huisstijl
BLAUW = "002664"        # BENU royal blue
GROEN = "58A618"        # BENU kelly green
GROEN_D = "407C05"      # BENU accessibility green
GROEN_L = "EEF9E5"      # BENU alabaster
LICHT = "EEF2F7"        # BENU anti-flash
LIJN = "D7DEE9"         # BENU general gray
GRIJS_L = "F5F7F9"      # BENU grey lightest
GRIJS = "546176"        # BENU black coral
BLAUW_M = "5C77A2"      # BENU blue yonder
BLAUW_P = "AFBED4"      # BENU pastel blue
ORANJE = "C25700"       # BENU accent orange (toegankelijk)
ORANJE_L = "FFF1E5"
WIT = "FFFFFF"

F = "Segoe UI"
FS = "Segoe UI Semibold"


def font(size=10, color=GRIJS, bold=False, italic=False, name=F):
    return Font(name=name, size=size, color=color, bold=bold, italic=italic)


def fill(kleur):
    return PatternFill("solid", fgColor=kleur)


def cf_fill(kleur):
    """Vulling voor voorwaardelijke opmaak (Excel gebruikt daar de achtergrondkleur)."""
    return PatternFill(fill_type="solid", start_color=kleur, end_color=kleur)


def side(kleur=LIJN, stijl="thin"):
    return Side(style=stijl, color=kleur)


f_norm = font()
f_label = font(color=BLAUW)
f_klein = font(9, GRIJS, italic=True)
f_invoer = font(10, BLAUW, bold=True)
f_kop_tabel = font(10, WIT, name=FS)
f_waarde = font(10, BLAUW)
f_sterk = font(10, BLAUW, name=FS)

rand_invoer = Border(left=side(BLAUW_M), right=side(BLAUW_M), top=side(BLAUW_M), bottom=side(BLAUW_M))
rand_cel = Border(bottom=side(LIJN))
links = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
midden = Alignment(horizontal="center", vertical="center", wrap_text=True)
rechts = Alignment(horizontal="right", vertical="center", indent=1)
boven = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)

wb = Workbook()
try:  # standaardstijl van de werkmap: Segoe UI 10
    wb._named_styles["Normal"].font = Font(name=F, size=10, color=GRIJS)
except Exception:
    pass


def kopbalk(sh, titel, sub, laatste_kol, hoogte_titel=40):
    """Blauwe kopbalk (rij 1-2) met groene accentlijn (rij 3) over kolommen A..laatste_kol."""
    kols = [chr(c) for c in range(ord("A"), ord(laatste_kol) + 1)]
    for c in kols:
        sh[f"{c}1"].fill = fill(BLAUW)
        sh[f"{c}2"].fill = fill(BLAUW)
        sh[f"{c}3"].fill = fill(GROEN)
    sh["B1"] = titel
    sh["B1"].font = Font(name=FS, size=20, color=WIT)
    sh["B1"].alignment = Alignment(vertical="bottom")
    sh["B2"] = sub
    sh["B2"].font = Font(name=F, size=10, color=BLAUW_P)
    sh["B2"].alignment = Alignment(vertical="top")
    sh.row_dimensions[1].height = hoogte_titel
    sh.row_dimensions[2].height = 22
    sh.row_dimensions[3].height = 4
    sh.sheet_view.showGridLines = False


def tabelkop(sh, rij, kol0, koppen, hoogte=30):
    for j, h in enumerate(koppen):
        c = sh.cell(row=rij, column=kol0 + j, value=h)
        c.font = f_kop_tabel
        c.fill = fill(BLAUW)
        c.alignment = midden
    sh.row_dimensions[rij].height = hoogte


def invoer(c, waarde=None, fmt=None):
    if waarde is not None:
        c.value = waarde
    c.font = f_invoer
    c.fill = fill(WIT)
    c.border = rand_invoer
    c.alignment = Alignment(horizontal="center", vertical="center")
    if fmt:
        c.number_format = fmt


# ================================================================== Omrekenfactoren
fs = wb.active
fs.title = "Omrekenfactoren"
ws = wb.create_sheet("Rotatie", 0)
for col, w in zip("ABCDEFGH", [2, 34, 13, 34, 13, 14, 44, 56]):
    fs.column_dimensions[col].width = w
kopbalk(fs, "Omrekenfactoren", "Volgens KNMP Informatorium Medicamentorum (Analgetica, G806). Witte velden zijn aanpasbaar; "
        "alle berekeningen gebruiken deze waarden.", "I")

fs["B5"] = "Instellingen"
fs["B5"].font = Font(name=FS, size=13, color=BLAUW)
fs.row_dimensions[5].height = 24
settings = [
    ("Verhouding morfine oraal : morfine s.c./i.v. (x : 1)", 3, "0.0",
     "KNMP IM Analgetica - opioïdrotatietabel (uit NVA-richtlijn Pijn bij kanker): 30 mg morfine oraal = 10 mg morfine s.c./i.v."),
    ("Dosis bij rotatie wegens bijwerkingen", 0.75, "0%",
     "KNMP IM Analgetica, Opioïdrotatie: bij rotatie vanwege bijwerkingen starten met 75% van de equivalente 24-uursdosering."),
    ("Dosis bij rotatie wegens onvoldoende pijnstilling", 1.0, "0%",
     "KNMP IM Analgetica: bij onvoldoende effect het nieuwe opioïd starten met een equivalente dosering."),
    ("Drempel 'hoge dosis' (OME mg/24 u) voor extra waarschuwing", 240, "0",
     "Afgeleid van KNMP: bij de meeste patiënten met kanker is < 240 mg morfine oraal/dag voldoende. Boven deze dosis extra alert."),
    ("Doorbraakbolus (fractie van dagdosis)", 1 / 6, "0.0%",
     "KNMP IM Analgetica, Doorbraakpijn: 1/6 van de (equivalente) dagdosis per keer, max. 4-6x per dag; "
     ">3x per dag nodig → onderhoudsdosering met 50% verhogen."),
    ("Doorbraakbolus - lagere variant (fractie van dagdosis)", 0.10, "0%",
     "Pallialine-pijnkaart: 10-15% van de dagdosis. Ter vergelijking; KNMP hanteert 1/6."),
    ("Beschikbare concentratie morfine 1 (mg/ml)", 10, "0.0#",
     "Eigen voorraad: morfine 10 mg/ml voor de pomp. Verschijnt in de keuzelijst op tabblad Rotatie."),
    ("Beschikbare concentratie morfine 2 (mg/ml)", 1, "0.0#",
     "Eigen voorraad: morfine 1 mg/ml, voor lage doseringen. Verschijnt in de keuzelijst op tabblad Rotatie."),
]
S0 = 6
for i, (lab, val, fmt, src) in enumerate(settings):
    r = S0 + i
    fs[f"B{r}"] = lab
    fs[f"B{r}"].font = f_label
    fs[f"B{r}"].alignment = links
    invoer(fs[f"C{r}"], val, fmt)
    fs[f"D{r}"] = src
    fs[f"D{r}"].font = f_klein
    fs[f"D{r}"].alignment = links
    fs.merge_cells(f"D{r}:H{r}")
    for c in "BDEFGH":
        fs[f"{c}{r}"].border = rand_cel
    fs.row_dimensions[r].height = 30
S_RATIO, S_SIDE, S_INSUF, S_HIGH, S_BOLLO, S_BOLHI, S_CONC1, S_CONC2 = [f"Omrekenfactoren!$C${S0 + i}" for i in range(8)]
CONC_LIJST = f"Omrekenfactoren!$C${S0 + 6}:$C${S0 + 7}"

fs["B15"] = "Omrekenfactoren naar OME (orale morfine-equivalenten, mg/24 u)"
fs["B15"].font = Font(name=FS, size=13, color=BLAUW)
fs.row_dimensions[15].height = 24
tabelkop(fs, 16, 2, ["Middel en toedieningsweg", "Eenheid invoer", "Equianalgetische verhouding",
                     "Factor → OME", "Max. dagdosis (invoereenheid)", "Bron", "Opmerking"], 34)

KNMP = "KNMP IM Analgetica (G806), Opioïdrotatie - conversiefactoren en rotatietabel (bron: NVA-richtlijn Pijn bij kanker)"
KNMPT = "KNMP IM Analgetica (G806), opioïdrotatietabel"
OTH = "Niet in KNMP-rotatietabel; gangbare literatuurwaarde"
ZMC = "Equivalentietabel opioïden, Zaans Medisch Centrum (2020)"
drugs = [
    ("Morfine oraal", "mg/24 u", "Referentie", 1, None, KNMP, ""),
    ("Morfine s.c./i.v.", "mg/24 u", "Morfine oraal : s.c./i.v. = 3 : 1 (zie instelling)", f"={S_RATIO}", None, KNMPT,
     "Volgt de instelling hierboven. KNMP-tabel: 60 mg oraal = 20 mg s.c./i.v."),
    ("Oxycodon oraal", "mg/24 u", "Morfine oraal : oxycodon oraal = 1,5 : 1", 1.5, None, KNMP,
     "KNMP-tabel: 60 mg morfine oraal = 40 mg oxycodon oraal."),
    ("Oxycodon s.c./i.v.", "mg/24 u", "Oxycodon oraal : s.c./i.v. = 2 : 1", 3, None, KNMPT,
     "KNMP-tabel: 60 mg morfine oraal = 20 mg oxycodon s.c./i.v."),
    ("Hydromorfon oraal", "mg/24 u", "Morfine oraal : hydromorfon oraal = 5 : 1", 5, None, KNMP,
     "KNMP-tabel: 60 mg morfine oraal = 12 mg hydromorfon oraal."),
    ("Hydromorfon s.c./i.v.", "mg/24 u", "Hydromorfon oraal : s.c./i.v. = 3 : 1", 15, None, KNMPT,
     "KNMP-tabel: 60 mg morfine oraal = 4 mg hydromorfon s.c./i.v."),
    ("Fentanyl pleister", "µg/uur", "Morfine oraal : fentanyl TD = 100 : 1", 2.4, None, KNMP,
     "µg/uur x 24 / 1000 x 100. KNMP-tabel: 25 µg/uur = 60 mg morfine oraal."),
    ("Fentanyl s.c./i.v. continu", "µg/uur", "Morfine parenteraal : fentanyl parenteraal = 100 : 1", 7.2, None,
     "KNMP IM Analgetica (G806): 'in een aantal bronnen 100:1 (beperkte gegevens)'",
     "µg/uur x 24 / 1000 x 100 x 3. LET OP: in de praktijk wordt ook 1:1 met de pleister (factor 2,4) gehanteerd; "
     "dat geeft een 3x lagere morfinedosis. Altijd overleg."),
    ("Fentanyl oromucosaal", "µg/24 u", "200 µg fentanyl oromucosaal/buccaal = 10 mg morfine oraal", 0.05, None,
     f"{ZMC}; bron: Perioperative Medicine for the Junior Clinician (Wiley, 2015)",
     "Zuig-, buccaal- of sublinguaaltablet (Actiq, Abstral, Effentora, Breakyl): totaal µg in 24 u x 0,05. "
     "Snelwerkend fentanyl wordt normaal los getitreerd; omrekening indicatief."),
    ("Fentanyl nasaal", "µg/24 u", "100 µg fentanyl nasaal = 30 mg morfine oraal", 0.3, None,
     f"{ZMC}; bron: EMA EPAR Instanyl (farmacokinetiek)",
     "Neusspray (Instanyl, PecFent): totaal µg in 24 u x 0,3. Omrekening indicatief."),
    ("Buprenorfine pleister", "µg/uur", "Morfine oraal : buprenorfine TD = 100 : 1", 2.4, 140, KNMP,
     "KNMP-tabel: 52,5 µg/uur = 120 mg morfine oraal. Boven 140 µg/uur niet geadviseerd (NVA)."),
    ("Tramadol oraal", "mg/24 u", "Morfine oraal : tramadol = 1 : 5 (KNMP-tabel)", 0.2, 400, KNMPT,
     "KNMP-tabel: 150 mg tramadol = 30 mg, 300 mg = 60 mg morfine oraal. (NVA COPD/hartfalen noemt 1 : 10.) Max. 400 mg/24 u."),
    ("Tapentadol oraal", "mg/24 u", "Morfine oraal : tapentadol = 1 : 2,5", 0.4, 500, KNMP,
     "KNMP-tabel: 150 mg tapentadol = 60 mg morfine oraal. Max. onderzochte dosis 500 mg/24 u."),
    ("Piritramide s.c./i.v.", "mg/24 u", "15-20 mg piritramide = 10 mg morfine parenteraal", 1.5, None,
     "KNMP IM Analgetica (G806), equivalentie piritramide",
     "Gerekend met 20 : 10 (voorzichtige kant bij rotatie naar morfine). Bij 15 : 10 wordt de factor 2,0."),
    ("Codeïne oraal", "mg/24 u", "Morfine oraal : codeïne = 1 : 10", 0.1, 240, OTH,
     "Niet in KNMP-rotatietabel; NHG raadt codeïne af als pijnstiller. Verifiëren."),
    ("Methadon oraal", "mg/24 u", "Dosisafhankelijk: 10 mg methadon = 40 / 60 / 80 mg morfine oraal", None, None,
     f"{ZMC}; bron: Pallialine, Opioïden. (KNMP: variabel 5 : 1 tot 10 : 1)",
     "Factor volgens de staffel hieronder (niet lineair). Altijd overleg met een deskundige "
     "(consultatieteam palliatieve zorg / pijnteam)."),
]
D0 = 17
for i, row in enumerate(drugs):
    r = D0 + i
    band = fill(GRIJS_L) if i % 2 else fill(WIT)
    for j, v in enumerate(row):
        c = fs.cell(row=r, column=2 + j, value=v)
        c.font = f_norm
        c.fill = band
        c.alignment = links
        c.border = rand_cel
    fs.cell(row=r, column=2).font = f_sterk
    fs.cell(row=r, column=3).alignment = midden
    fs.cell(row=r, column=7).font = f_klein
    fs.cell(row=r, column=8).font = font(9, GRIJS)
    for col in (5, 6):
        c = fs.cell(row=r, column=col)
        is_formule = isinstance(c.value, str) and c.value.startswith("=")
        invoer(c, fmt="0.0##" if col == 5 else "0")
        if is_formule:
            c.font = f_waarde
            c.fill = fill(LICHT)
    fs.row_dimensions[r].height = 32
D1 = D0 + len(drugs) - 1
NAMES = f"Omrekenfactoren!$B${D0}:$B${D1}"
UNITS = f"Omrekenfactoren!$C${D0}:$C${D1}"
FACT = f"Omrekenfactoren!$E${D0}:$E${D1}"
MAXD = f"Omrekenfactoren!$F${D0}:$F${D1}"

# Methadon: factor hangt af van de methadondosis (staffel, niet lineair)
M0 = D1 + 3
fs[f"B{M0 - 1}"] = "Methadon oraal - dosisafhankelijke factor"
fs[f"B{M0 - 1}"].font = Font(name=FS, size=13, color=BLAUW)
fs.row_dimensions[M0 - 1].height = 24
tabelkop(fs, M0, 2, ["Methadon (mg/24 u)", "Tot en met (mg/24 u)", "10 mg methadon =", "Factor → OME", "", "Bron"], 30)
staffel = [("0-20 mg", 20, "40 mg morfine oraal", 4), ("21-50 mg", 50, "60 mg morfine oraal", 6),
           ("> 50 mg", None, "80 mg morfine oraal", 8)]
for i, (lab, tm, eq, f) in enumerate(staffel):
    r = M0 + 1 + i
    for j, v in enumerate((lab, tm, eq, f, None, f"{ZMC}; bron: Pallialine, Opioïden")):
        c = fs.cell(row=r, column=2 + j, value=v)
        c.font = f_norm
        c.alignment = links
        c.border = rand_cel
    fs.cell(row=r, column=2).font = f_sterk
    fs.cell(row=r, column=7).font = f_klein
    if tm is not None:
        invoer(fs.cell(row=r, column=3), tm, "0")
    invoer(fs.cell(row=r, column=5), f, "0.0#")
    fs.row_dimensions[r].height = 22
fs[f"B{M0 + 4}"] = ("Omrekenfactor niet lineair: bij een methadondosis rond een grens verandert de OME sprongsgewijs. "
                    "Altijd overleg met een deskundige.")
fs[f"B{M0 + 4}"].font = f_klein
fs.merge_cells(f"B{M0 + 4}:H{M0 + 4}")
M_TM1, M_TM2 = f"Omrekenfactoren!$C${M0 + 1}", f"Omrekenfactoren!$C${M0 + 2}"
M_F1, M_F2, M_F3 = (f"Omrekenfactoren!$E${M0 + 1 + i}" for i in range(3))


def methadon_factor(d):
    """Excel-formule (zonder '=') voor de methadonfactor bij dosis d (mg/24 u)."""
    return f"IF({d}<={M_TM1},{M_F1},IF({d}<={M_TM2},{M_F2},{M_F3}))"


def methadon_dosis(ome):
    """Excel-formule (zonder '='): methadondosis bij OME (staffel omgekeerd; tussen twee stappen de lagere dosis)."""
    return (f"IF({ome}<={M_TM1}*{M_F1},{ome}/{M_F1},IF({ome}<={M_TM2}*{M_F2},MAX({M_TM1},{ome}/{M_F2}),"
            f"MAX({M_TM2},{ome}/{M_F3})))")

fs.freeze_panes = "A4"
fs.page_setup.orientation = "landscape"
fs.page_setup.fitToWidth = 1
fs.page_setup.fitToHeight = 0
fs.sheet_properties.pageSetUpPr.fitToPage = True
fs.sheet_properties.tabColor = GROEN

# ================================================================== Rotatie
for col, w in zip("ABCDEFGHIJ", [2, 15, 28, 23, 11, 12, 14, 30, 56, 2]):
    ws.column_dimensions[col].width = w
kopbalk(ws, "Opioïdrotatie", "naar morfine s.c./i.v. continu (pomp)   ·   omrekening volgens KNMP Informatorium Medicamentorum",
        "J", 44)
ws.sheet_view.showRowColHeaders = False
ws.sheet_view.zoomScale = 100
ws.sheet_properties.tabColor = BLAUW

# Patiëntregel (B4; wordt door de macro gevuld)
PLAATS = "Nog niet uitgelezen uit Pharmacom  ·  open het dossier en druk op Ctrl+Alt+F3"
ws.merge_cells("B4:I4")
ws["B4"] = PLAATS
ws["B4"].font = Font(name=FS, size=10, color=GROEN_D)
ws["B4"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
for c in "BCDEFGHI":
    ws[f"{c}4"].fill = fill(GROEN_L)
ws.row_dimensions[4].height = 26

# Legenda
ws["B5"] = "Legenda"
ws["B5"].font = f_klein
ws["B5"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
for ref, tekst, f_, fi, br in (("C5", "Invulveld", f_invoer, WIT, rand_invoer),
                               ("D5", "Berekend", f_waarde, LICHT, None),
                               ("E5", "Let op", font(9, ORANJE, name=FS), ORANJE_L, None)):
    c = ws[ref]
    c.value = tekst
    c.font = Font(name=f_.name, size=9, color=f_.color, bold=f_.bold)
    c.fill = fill(fi)
    c.alignment = midden
    if br:
        c.border = br
ws.merge_cells("E5:F5")
ws.row_dimensions[5].height = 20
ws.row_dimensions[6].height = 8


def stap(r, nummer, titel, sub=None):
    """Kop van een stap: groen nummer + blauwe titel, groene lijn eronder."""
    c = ws[f"B{r}"]
    c.value = nummer
    c.font = Font(name=FS, size=12, color=WIT)
    c.fill = fill(GROEN)
    c.alignment = midden
    ws[f"C{r}"] = titel
    ws[f"C{r}"].font = Font(name=FS, size=13, color=BLAUW)
    ws[f"C{r}"].alignment = Alignment(vertical="center", indent=1)
    for k in "BCDEFGHI":
        ws[f"{k}{r}"].border = Border(bottom=side(GROEN, "medium"))
    ws.row_dimensions[r].height = 26
    if sub:
        ws[f"C{r + 1}"] = sub
        ws[f"C{r + 1}"].font = f_klein
        ws[f"C{r + 1}"].alignment = Alignment(vertical="center", indent=1)
        ws.merge_cells(f"C{r + 1}:I{r + 1}")
        ws.row_dimensions[r + 1].height = 20


def label(r, tekst, noot=None):
    ws[f"B{r}"] = tekst
    ws[f"B{r}"].font = f_label
    ws[f"B{r}"].alignment = links
    ws.merge_cells(f"B{r}:C{r}")
    if noot:
        ws[f"E{r}"] = noot
        ws[f"E{r}"].font = f_klein
        ws[f"E{r}"].alignment = links
        ws.merge_cells(f"E{r}:I{r}")
    ws.row_dimensions[r].height = 22


def berekend(ref, formule, fmt, sterk=False):
    c = ws[ref]
    c.value = formule
    c.number_format = fmt
    c.font = f_sterk if sterk else f_waarde
    c.fill = fill(LICHT)
    c.alignment = rechts


# ---------------------------------------------------------------- 1. huidig gebruik
stap(7, 1, "Huidig opioïdgebruik",
     "Alle opioïden van de afgelopen 24 uur, inclusief gebruikte doorbraakmedicatie. Pleisters en continue fentanyl in µg/uur.")
tabelkop(ws, 9, 2, ["Soort", "Middel + toedieningsweg", "Dosis per 24 u", "Eenheid", "Factor → OME", "OME (mg/24 u)",
                    "Controle", "Uitgelezen uit Pharmacom"], 30)
R0, R1 = 10, 15
dv_soort = DataValidation(type="list", formula1='"Onderhoud,Doorbraak"', allow_blank=True)
dv_mid = DataValidation(type="list", formula1=NAMES, allow_blank=True,
                        error="Kies een middel uit de lijst (tabblad Omrekenfactoren).", showErrorMessage=True)
dv_num = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                        error="Voer een getal ≥ 0 in.", showErrorMessage=True)
for v in (dv_soort, dv_mid, dv_num):
    ws.add_data_validation(v)
for r in range(R0, R1 + 1):
    invoer(ws[f"B{r}"])
    invoer(ws[f"C{r}"])
    ws[f"C{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    invoer(ws[f"D{r}"], fmt="0.0##")
    dv_soort.add(f"B{r}")
    dv_mid.add(f"C{r}")
    dv_num.add(f"D{r}")
    berekend(f"E{r}", f'=IF(C{r}="","",INDEX({UNITS},MATCH(C{r},{NAMES},0)))', "@")
    ws[f"E{r}"].alignment = midden
    berekend(f"F{r}", f'=IF(C{r}="","",IF(C{r}="Methadon oraal",IF(D{r}="","",{methadon_factor(f"D{r}")}),'
                      f'IF(INDEX({FACT},MATCH(C{r},{NAMES},0))="","n.v.t.",INDEX({FACT},MATCH(C{r},{NAMES},0)))))',
             "0.0##")
    berekend(f"G{r}", f'=IF(OR(C{r}="",D{r}=""),"",IF(ISNUMBER(F{r}),D{r}*F{r},"overleg"))', "0.0", sterk=True)
    c = ws[f"H{r}"]
    c.value = (f'=IF(OR(C{r}="",D{r}=""),"",IF(NOT(ISNUMBER(F{r})),"Geen omrekening - overleg deskundige",'
               f'IF(C{r}="Methadon oraal","Dosisafhankelijke factor - overleg",'
               f'IF(C{r}="Fentanyl oromucosaal","Indicatieve factor - controleer",'
               f'IF(C{r}="Fentanyl nasaal","Indicatieve factor - controleer",'
               f'IF(AND(ISNUMBER(INDEX({MAXD},MATCH(C{r},{NAMES},0))),'
               f'D{r}>INDEX({MAXD},MATCH(C{r},{NAMES},0))),'
               f'"Boven max. dagdosis ("&INDEX({MAXD},MATCH(C{r},{NAMES},0))&")","✓ OK"))))))')
    c.font = font(9, GROEN_D, name=FS)
    c.fill = fill(LICHT)
    c.alignment = links
    c = ws[f"I{r}"]
    c.font = font(8.5, GRIJS)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    for k in "BCDEFGHI":
        ws[f"{k}{r}"].border = Border(left=ws[f"{k}{r}"].border.left, right=ws[f"{k}{r}"].border.right,
                                      top=ws[f"{k}{r}"].border.top, bottom=side(LIJN) if k in "EFGHI" else side(BLAUW_M))
    ws.row_dimensions[r].height = 34
# Controle in oranje als er iets mis is
ws.conditional_formatting.add(f"H{R0}:H{R1}", FormulaRule(formula=[f'AND(H{R0}<>"",H{R0}<>"✓ OK")'],
                                                          font=Font(color=ORANJE, bold=True), fill=cf_fill(ORANJE_L)))
# voorbeeld
ws[f"B{R0}"], ws[f"C{R0}"], ws[f"D{R0}"] = "Onderhoud", "Oxycodon oraal", 80
ws[f"B{R0 + 1}"], ws[f"C{R0 + 1}"], ws[f"D{R0 + 1}"] = "Doorbraak", "Oxycodon oraal", 30
ws[f"C{R0}"].comment = Comment("Voorbeeld: oxycodon retard 2dd 40 mg = 80 mg/24 u. Overschrijven of 'Wis invoer'.", "Rekenhulp")
ws[f"C{R0 + 1}"].comment = Comment("Voorbeeld: 3x oxycodon IR 10 mg in afgelopen 24 u = 30 mg.", "Rekenhulp")

TR = R1 + 1
ws[f"B{TR}"] = "Totaal orale morfine-equivalenten (OME)"
ws[f"B{TR}"].font = f_sterk
ws[f"B{TR}"].alignment = links
ws.merge_cells(f"B{TR}:F{TR}")
berekend(f"G{TR}", f"=SUM(G{R0}:G{R1})", "0.0", sterk=True)
ws[f"G{TR}"].font = Font(name=FS, size=11, color=BLAUW)
ws[f"H{TR}"] = "mg morfine oraal / 24 u"
ws[f"H{TR}"].font = f_klein
ws[f"H{TR}"].alignment = links
for k in "BCDEFGHI":
    ws[f"{k}{TR}"].border = Border(top=side(BLAUW, "medium"))
ws.row_dimensions[TR].height = 26
ws.row_dimensions[TR + 1].height = 12

# ---------------------------------------------------------------- 2. reden en patiënt
stap(18, 2, "Reden van rotatie en patiënt")
reasons = ["Bijwerkingen", "Onvoldoende pijnstilling", "Anders (handmatig %)"]
dv_reason = DataValidation(type="list", formula1='"' + ",".join(reasons) + '"', allow_blank=False)
ws.add_data_validation(dv_reason)
label(19, "Reden van rotatie", "Onvoldoende pijnstilling → 100%; bijwerkingen → 75% (KNMP; instelbaar op tabblad Omrekenfactoren).")
invoer(ws["D19"], "Onvoldoende pijnstilling")
dv_reason.add("D19")
label(20, "Handmatig % van equianalgetische dosis", "Alleen gebruikt bij 'Anders (handmatig %)'.")
invoer(ws["D20"], None, "0%")
label(21, "Extra reductie (optioneel)", "Bijv. 25% bij een kwetsbare of oudere patiënt of bij sufheid. Komt bovenop de reden.")
invoer(ws["D21"], 0, "0%")
label(22, "eGFR (ml/min/1,73 m²)", "Uit Pharmacom (Diagnostisch dossier) of zelf invullen. Geeft een waarschuwing bij nierfunctiestoornis.")
invoer(ws["D22"], None, "0")
pct = DataValidation(type="decimal", operator="between", formula1="0", formula2="1.5", allow_blank=True,
                     error="Voer een percentage in (bijv. 75%).", showErrorMessage=True)
ws.add_data_validation(pct)
pct.add("D20")
pct.add("D21")
label(23, "Toe te passen % van de equianalgetische dosis")
berekend("D23", f'=IF(D19="Bijwerkingen",{S_SIDE},IF(D19="Onvoldoende pijnstilling",{S_INSUF},'
                f'IF(ISNUMBER(D20),D20,0)))*(1-N(D21))', "0%", sterk=True)
ws["B23"].font = f_sterk
ws.row_dimensions[24].height = 12

# ---------------------------------------------------------------- 3. pomp
stap(25, 3, "Pomp")
label(26, "Concentratie morfine in spuit/cassette (mg/ml)",
      "Keuzelijst: 10 of 1 mg/ml (eigen voorraad; instelbaar op tabblad Omrekenfactoren).")
invoer(ws["D26"], 10, "0.0#")
dv_conc = DataValidation(type="list", formula1=CONC_LIJST, allow_blank=False,
                         error="Kies een concentratie uit de lijst (tabblad Omrekenfactoren).", showErrorMessage=True)
ws.add_data_validation(dv_conc)
dv_conc.add("D26")
label(27, "Kleinste instelstap van de pomp (ml/uur)", "Bijv. 0,01 of 0,1 - afhankelijk van het type pomp.")
invoer(ws["D27"], 0.01, "0.0##")
label(28, "Volume spuit/cassette (ml)", "Voor de looptijd.")
invoer(ws["D28"], 50, "0")
for ref in ("D27", "D28"):
    dv_num.add(ref)
ws.row_dimensions[29].height = 12

# ---------------------------------------------------------------- 4. resultaat: tegels
stap(30, 4, "Resultaat: morfine s.c./i.v. continu")
ws.row_dimensions[31].height = 8
# Detailregels (onder de tegels), eerst vastleggen voor de verwijzingen
DET = 41
det = [
    ("Totaal OME", f"=G{TR}", "0.0", "mg morfine oraal / 24 u"),
    ("Equianalgetisch: morfine s.c./i.v.", f"=D{DET}/{S_RATIO}", "0.0", "mg / 24 u  (OME ÷ verhouding oraal : s.c.)"),
    ("Toegepast percentage", "=D23", "0%", "reden van rotatie × (1 − extra reductie)"),
    ("Nieuwe dagdosis morfine s.c./i.v.", f"=D{DET + 1}*D{DET + 2}", "0.0", "mg / 24 u"),
    ("Dosis per uur", f"=D{DET + 3}/24", "0.00", "mg / uur"),
    ("Inloopsnelheid (exact)", f"=IF(N(D26)>0,D{DET + 4}/D26,0)", "0.000", "ml / uur"),
    ("Inloopsnelheid afgerond op pompstap", f"=IF(N(D27)>0,ROUND(D{DET + 5}/D27,0)*D27,D{DET + 5})", "0.00",
     "ml / uur  ← instellen op de pomp"),
    ("Werkelijke dagdosis bij afgeronde snelheid", f"=D{DET + 6}*D26*24", "0.0", "mg / 24 u"),
    ("Afwijking door afronding", f"=IF(D{DET + 3}>0,D{DET + 7}/D{DET + 3}-1,0)", "+0.0%;-0.0%;0.0%",
     "t.o.v. de berekende dagdosis"),
    ("Volume per 24 uur", f"=D{DET + 6}*24", "0.0", "ml / 24 u"),
    ("Looptijd spuit/cassette", f"=IF(D{DET + 6}>0,D28/D{DET + 6},0)", "0", "uur"),
    ("Doorbraakbolus s.c. - KNMP (1/6 dagdosis)", f"=D{DET + 7}*{S_BOLLO}", "0.0", "mg per keer; max. 4-6x per dag (KNMP)"),
    ("Doorbraakbolus s.c. - lagere variant (10%)", f"=D{DET + 7}*{S_BOLHI}", "0.0",
     "mg per keer (Pallialine 10-15%), bijv. bij een kwetsbare patiënt"),
    ("Doorbraakbolus in ml (KNMP  |  lagere variant)", None, "@", "ml per bolus; interval volgens lokaal protocol"),
    ("Andere concentratie (ter vergelijking)", f"=IF(N(D26)={S_CONC1},{S_CONC2},{S_CONC1})", '0.0#" mg/ml"',
     "de andere sterkte uit de voorraad"),
    ("Pompstand bij andere concentratie",
     f"=IF(AND(N(D{DET + 14})>0,D{DET + 3}>0),"
     f"IF(N(D27)>0,ROUND(D{DET + 4}/D{DET + 14}/D27,0)*D27,D{DET + 4}/D{DET + 14}),0)", "0.00",
     f'=IF(D{DET + 15}>0,"ml / uur  ·  afwijking "&FIXED((D{DET + 15}*D{DET + 14}*24/D{DET + 3}-1)*100,1,TRUE)'
     f'&"%  ·  looptijd "&FIXED(D28/D{DET + 15},0,TRUE)&" uur","ml / uur")'),
]
(R_OME, R_SC, R_PCT, R_DAG, R_UUR, R_EXACT, R_ML, R_ECHT, R_AFW, R_VOL, R_LOOP, R_BOL, R_BOL2, R_BOLML,
 R_ALTC, R_ALTML) = range(DET, DET + 16)

# Tegels (rij 32-36)
tegels = [
    # (kolommen, label, waarde, formaat, subtekst, hoofdtegel)
    ("B", "C", "POMPSTAND", f'=IF(G{TR}>0,D{R_ML},"")', '0.00" ml/uur"',
     f'=IF(G{TR}>0,"bij "&FIXED(D26,1,TRUE)&" mg/ml  ·  stap "&FIXED(D27,2,TRUE)&" ml/uur","Vul stap 1 in")', True),
    ("D", "F", "MORFINE S.C./I.V.", f'=IF(G{TR}>0,D{R_ECHT},"")', '0.0" mg/24 u"',
     f'=IF(G{TR}>0,"= "&FIXED(D{R_ECHT}/24,2,TRUE)&" mg/uur  ·  "&TEXT(D23,"0%")&" van equianalgetisch","")', False),
    ("G", "H", "DOORBRAAKBOLUS", f'=IF(G{TR}>0,D{R_BOL},"")', '0.0" mg s.c."',
     f'=IF(AND(G{TR}>0,N(D26)>0),"= "&FIXED(D{R_BOL}/D26,2,TRUE)&" ml  ·  1/6 dagdosis, max. 4-6× per dag","")', False),
    ("I", "I", "SPUIT / CASSETTE", f'=IF(G{TR}>0,D{R_LOOP},"")', '0" uur looptijd"',
     f'=IF(G{TR}>0,FIXED(D{R_VOL},1,TRUE)&" ml per 24 u  ·  spuit van "&D28&" ml","")', False),
]
for k0, k1, lab, waarde, fmt, sub, hoofd in tegels:
    achter = BLAUW if hoofd else LICHT
    voor = WIT if hoofd else BLAUW
    subkleur = BLAUW_P if hoofd else GRIJS
    for r in range(32, 37):
        for k in range(ord(k0), ord(k1) + 1):
            ws[f"{chr(k)}{r}"].fill = fill(achter)
    ws.merge_cells(f"{k0}32:{k1}32")
    ws.merge_cells(f"{k0}33:{k1}34")
    ws.merge_cells(f"{k0}35:{k1}35")
    ws[f"{k0}32"] = lab
    ws[f"{k0}32"].font = Font(name=FS, size=9, color=GROEN if hoofd else GROEN_D)
    ws[f"{k0}32"].alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
    ws[f"{k0}33"] = waarde
    ws[f"{k0}33"].number_format = fmt
    ws[f"{k0}33"].font = Font(name=FS, size=22 if hoofd else 18, color=voor)
    ws[f"{k0}33"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws[f"{k0}35"] = sub
    ws[f"{k0}35"].font = Font(name=F, size=8.5, color=subkleur)
    ws[f"{k0}35"].alignment = Alignment(horizontal="left", vertical="top", indent=1, wrap_text=True)
for r, h in ((32, 20), (33, 22), (34, 22), (35, 30), (36, 6)):
    ws.row_dimensions[r].height = h
# Witte scheiding tussen de tegels
for k0, *_ in tegels[1:]:
    for r in range(32, 37):
        ws[f"{k0}{r}"].border = Border(left=Side(style="thick", color=WIT))

# Altijd-blok direct onder de tegels
ws.merge_cells("B37:I38")
ws["B37"] = ("Altijd:  start de pomp op het moment van de volgende geplande dosis van het oude middel  ·  zorg voor doorbraakmedicatie "
             "(KNMP: 1/6 dagdosis, max. 4-6× per dag; > 3× per dag nodig → onderhoud +50%)  ·  evalueer pijn, sufheid en ademhaling  ·  "
             "laat berekening en pompinstelling dubbel controleren.")
ws["B37"].font = font(9, BLAUW)
ws["B37"].alignment = Alignment(wrap_text=True, vertical="center", indent=1)
for rr in (37, 38):
    ws[f"B{rr}"].fill = fill(GROEN_L)
    for k in "CDEFGHI":
        ws[f"{k}{rr}"].fill = fill(GROEN_L)
    ws.row_dimensions[rr].height = 18
ws.row_dimensions[39].height = 12

# Detailtabel
ws["B40"] = "Berekening in detail"
ws["B40"].font = Font(name=FS, size=10, color=BLAUW)
ws["B40"].alignment = Alignment(vertical="center", indent=1)
ws.merge_cells("B40:C40")
for k in "BCDEFGHI":
    ws[f"{k}40"].border = Border(bottom=side(LIJN))
ws.row_dimensions[40].height = 22
for i, (lab, fm, fmt, uitleg) in enumerate(det):
    r = DET + i
    label(r, lab, uitleg)
    ws.row_dimensions[r].height = 20
    if fm:
        berekend(f"D{r}", fm, fmt, sterk=r in (R_DAG, R_ML))
    for k in "BCDEFGHI":
        ws[f"{k}{r}"].border = Border(bottom=side(GRIJS_L))
berekend(f"D{R_BOLML}", f'=IF(N(D26)>0,FIXED(D{R_BOL}/D26,2,TRUE)&"  |  "&FIXED(D{R_BOL2}/D26,2,TRUE),"")', "@")
for r in (R_DAG, R_ML):
    ws[f"B{r}"].font = f_sterk
    ws[f"D{r}"].fill = fill(GROEN_L)
    ws[f"D{r}"].font = Font(name=FS, size=10, color=GROEN_D)
R_DET_END = DET + len(det) - 1
ws.row_dimensions[R_DET_END + 1].height = 12

# ---------------------------------------------------------------- 5. controles
CW = R_DET_END + 2
stap(CW, 5, "Controles en waarschuwingen",
     f"Rekenhulp - vervangt geen klinisch oordeel, lokaal protocol of dubbele controle.  Bron: KNMP Informatorium Medicamentorum, "
     f"Analgetica (G806).  {VERSIE}")
ws[f"C{CW + 1}"].font = font(8.5, BLAUW_M, italic=True)
L = f"C{R0}:C{R1}"
warns = [
    f'=IF(COUNTIF({L},"Methadon*")>0,"Methadon: meegerekend met een dosisafhankelijke factor (10 mg = 40/60/80 mg morfine oraal; '
    f'tabblad Omrekenfactoren). Niet lineair, lange en variabele halfwaardetijd: overleg altijd met een deskundige.","")',
    f'=IF(COUNTIF({L},"Fentanyl oromucosaal")+COUNTIF({L},"Fentanyl nasaal")>0,"Snelwerkend fentanyl (oromucosaal/nasaal): '
    f'de omrekenfactor is indicatief (geen vaste equivalentie; wordt normaal los getitreerd). Vul het werkelijke gebruik in µg/24 u in '
    f'en controleer de uitkomst.","")',
    f'=IF(COUNTIF({L},"Fentanyl pleister")>0,"Fentanylpleister: na verwijderen wordt nog 12-24 uur fentanyl afgegeven. Start de pomp in overleg (bijv. ca. 12 uur na verwijderen of met een lagere startdosis) en zorg voor doorbraakmedicatie.","")',
    f'=IF(COUNTIF({L},"Fentanyl s.c./i.v. continu")>0,"Fentanyl s.c./i.v.: de KNMP-factor parenteraal morfine : fentanyl 100 : 1 berust op beperkte gegevens. Bij 1 : 1 met de pleister zou de morfinedosis 3x lager uitkomen. Altijd overleg.","")',
    f'=IF(COUNTIF({L},"Buprenorfine pleister")>0,"Buprenorfinepleister: partiële agonist met lange werkingsduur. Plan de overgang naar morfine in overleg.","")',
    '=IF(AND(ISNUMBER(D22),D22<30),"eGFR < 30: morfine en metabolieten stapelen. Morfine is minder geschikt - overweeg een ander opioïd (bijv. fentanyl of hydromorfon) of een fors lagere dosis in overleg.","")',
    f'=IF(G{TR}>={S_HIGH},"Hoge dosis (OME ≥ "&{S_HIGH}&" mg/24 u): de omrekening is onzekerder. Overweeg extra reductie, stapsgewijze rotatie en overleg met het consultatieteam palliatieve zorg of pijnteam.","")',
    f'=IF(COUNTIF(H{R0}:H{R1},"Boven*")>0,"Een of meer doseringen liggen boven de maximale dagdosis - controleer de invoer.","")',
    f'=IF(COUNTIFS(B{R0}:B{R1},"Doorbraak",I{R0}:I{R1},"*MAXIMUM*")>0,"Zo nodig / bij pijn is uit Pharmacom met het MAXIMUM ingevuld. Verlaag naar het werkelijke gebruik als dat minder was; anders wordt de pompdosis te hoog.","")',
    '=IF(AND(D19="Onvoldoende pijnstilling",N(D21)=0),"Volledige equianalgetische dosis (geen reductie). Evalueer effect en bijwerkingen na enkele uren.","")',
    f'=IF(ABS(D{R_AFW})>0.1,"Afronding op de pompstap wijkt "&FIXED(ABS(D{R_AFW})*100,0,TRUE)&"% af bij "'
    f'&FIXED(D26,1,TRUE)&" mg/ml."&IF(AND(D{R_ALTML}>0,ABS(D{R_ALTML}*D{R_ALTC}*24/D{R_DAG}-1)<=0.1),'
    f'" Kies "&FIXED(D{R_ALTC},0,TRUE)&" mg/ml: pompstand "&FIXED(D{R_ALTML},2,TRUE)&" ml/uur.",'
    f'" Overweeg een andere concentratie of pompstap."),"")',
    f'=IF(AND(G{TR}>0,D{R_LOOP}>0,D{R_LOOP}<24),"De spuit/cassette is binnen 24 uur leeg ("&FIXED(D{R_LOOP},0,TRUE)'
    f'&" uur bij "&FIXED(D26,1,TRUE)&" mg/ml). Overweeg "&FIXED(MAX({S_CONC1},{S_CONC2}),0,TRUE)&" mg/ml of een groter volume.","")',
    '=IF(AND(D19="Anders (handmatig %)",NOT(ISNUMBER(D20))),"Vul bij \'Anders\' een handmatig percentage in.","")',
]
# Verborgen hulpkolommen L (tekst) en M (volgnummer van niet-lege teksten); de zichtbare
# regels tonen alleen de echte waarschuwingen, bovenaan, zonder lege tussenrijen.
W0 = CW + 2
W1 = W0 + len(warns) - 1
for i, f in enumerate(warns):
    r = W0 + i
    ws[f"L{r}"] = f
    ws[f"M{r}"] = f'=IF(L{r}<>"",COUNTIF(L${W0}:L{r},"?*"),"")'
ws.column_dimensions["L"].hidden = True
ws.column_dimensions["M"].hidden = True
AANTAL = f'COUNTIF($L${W0}:$L${W1},"?*")'
for k in range(1, len(warns) + 1):
    r = W0 + k - 1
    ws.merge_cells(f"C{r}:I{r}")
    tekst = f'IFERROR(INDEX($L${W0}:$L${W1},MATCH({k},$M${W0}:$M${W1},0)),"")'
    if k == 1:
        ws[f"C{r}"] = f'=IF({AANTAL}=0,"Geen bijzonderheden gevonden.",{tekst})'
        ws[f"B{r}"] = f'=IF({AANTAL}=0,"✓","⚠")'
    else:
        ws[f"C{r}"] = f"={tekst}"
        ws[f"B{r}"] = f'=IF(LEN(C{r})>0,"⚠","")'
    ws[f"C{r}"].font = font(9.5, ORANJE, name=FS)
    ws[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    ws[f"B{r}"].font = Font(name="Segoe UI Symbol", size=11, color=ORANJE)
    ws[f"B{r}"].alignment = midden
    ws.row_dimensions[r].height = 30
    ws.conditional_formatting.add(f"B{r}:I{r}", FormulaRule(formula=[f'$B${r}="⚠"'], fill=cf_fill(ORANJE_L)))
ws.conditional_formatting.add(f"B{W0}:I{W0}", FormulaRule(formula=[f'$B${W0}="✓"'], fill=cf_fill(GROEN_L),
                                                          font=Font(color=GROEN_D)))
EIND = W1
ws.freeze_panes = "A5"
ws.print_area = f"A1:J{EIND}"
ws.print_options.horizontalCentered = True
ws.page_setup.orientation = "portrait"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins.left = ws.page_margins.right = 0.4
ws.page_margins.top = ws.page_margins.bottom = 0.5
ws.oddFooter.left.text = "Opioïdrotatie - rekenhulp"
ws.oddFooter.left.size = 8
ws.oddFooter.center.text = "Pagina &P van &N"
ws.oddFooter.center.size = 8
ws.oddFooter.right.text = "&D &T"
ws.oddFooter.right.size = 8

# ================================================================== Omrekentabel
qs = wb.create_sheet("Omrekentabel")
cols = [n for n, *_ in drugs]
extra = ["Morfine pomp mg/uur", "Morfine pomp ml/uur"]
laatste = chr(ord("B") + len(cols) + len(extra))
kopbalk(qs, "Snelle omrekentabel", "Equianalgetisch, zonder reductie. Bij rotatie wegens bijwerkingen: 75% van deze waarden. "
        "Pompkolommen gebruiken de concentratie van tabblad Rotatie.", chr(ord(laatste) + 1))
qs.column_dimensions["A"].width = 2
qs.column_dimensions["B"].width = 16
for j in range(len(cols) + len(extra)):
    qs.column_dimensions[chr(ord("C") + j)].width = 12.5
tabelkop(qs, 5, 2, ["OME\n(mg morfine oraal / 24 u)"] + cols + extra, 46)
qs.cell(row=6, column=2, value="").fill = fill(LICHT)
for j, n in enumerate(cols):
    c = qs.cell(row=6, column=3 + j, value=f'=INDEX({UNITS},MATCH("{n}",{NAMES},0))')
    c.font = f_klein
    c.fill = fill(LICHT)
    c.alignment = midden
for j, t in enumerate(("mg/uur", "ml/uur")):
    c = qs.cell(row=6, column=3 + len(cols) + j, value=t)
    c.font = f_klein
    c.fill = fill(LICHT)
    c.alignment = midden
for i, ome in enumerate([30, 60, 90, 120, 180, 240, 300, 360, 480, 600]):
    r = 7 + i
    band = fill(GRIJS_L) if i % 2 else fill(WIT)
    invoer(qs.cell(row=r, column=2), ome)
    for j, n in enumerate(cols):
        if n == "Methadon oraal":
            v = "=" + methadon_dosis(f"$B{r}")
        else:
            v = f'=$B{r}/INDEX({FACT},MATCH("{n}",{NAMES},0))'
        c = qs.cell(row=r, column=3 + j, value=v)
        c.number_format = "0.0"
        c.font = f_waarde
        c.fill = band
        c.alignment = midden
        c.border = rand_cel
    m = qs.cell(row=r, column=3 + len(cols), value=f"=$B{r}/{S_RATIO}/24")
    m.number_format = "0.00"
    l_ = qs.cell(row=r, column=4 + len(cols), value=f"=IF(N(Rotatie!$D$26)>0,$B{r}/{S_RATIO}/24/Rotatie!$D$26,0)")
    l_.number_format = "0.000"
    for c in (m, l_):
        c.font = Font(name=FS, size=10, color=GROEN_D)
        c.fill = fill(GROEN_L)
        c.alignment = midden
        c.border = rand_cel
    qs.row_dimensions[r].height = 20
qs["B18"] = "Witte OME-waarden in kolom B zijn aanpasbaar."
qs["B18"].font = f_klein
qs["B19"] = ("Methadon: dosisafhankelijke staffel (tabblad Omrekenfactoren); valt de OME tussen twee stappen, dan staat "
             "hier de lagere methadondosis. Fentanyl oromucosaal/nasaal: indicatief. Altijd overleg.")
qs["B19"].font = f_klein
qs.freeze_panes = "C7"
qs.sheet_properties.tabColor = BLAUW_M
qs.page_setup.orientation = "landscape"
qs.page_setup.fitToWidth = 1
qs.page_setup.fitToHeight = 1
qs.sheet_properties.pageSetUpPr.fitToPage = True

# ================================================================== Doseercodes
ds = wb.create_sheet("Doseercodes", 2)
for col, w in zip("ABCDEFG", [2, 14, 20, 15, 20, 72, 2]):
    ds.column_dimensions[col].width = w
kopbalk(ds, "Eigen doseercodes", "Pharmacom-doseercodes die de rekenhulp niet zelf herkent. Een code hier gaat altijd vóór de "
        "automatische vertaling.", "G")
uitleg = [
    "De rekenhulp vertaalt doseercodes volgens de Centrale doseertabel (Health Base): frequentie - tijdsaanduiding - "
    "hoeveelheid - toedieningsvorm, eventueel 'M' + getal = niet meer dan ... per dag (bijv. P1CM6 = bij pijn 1 capsule, max. 6 per dag).",
    "Eenheden per 24 u = aantal tabletten/capsules/ml per dag (bij een pleister: aantal tegelijk geplakte pleisters). "
    "Zo nodig = J voor zo nodig/bij pijn: dan wordt het maximum per 24 u ingevuld; verlaag zelf naar het werkelijke gebruik.",
]
for i, t in enumerate(uitleg):
    c = ds.cell(row=5 + i, column=2, value=t)
    c.font = font(9, GRIJS)
    c.alignment = boven
    ds.merge_cells(start_row=5 + i, start_column=2, end_row=5 + i, end_column=6)
    ds.row_dimensions[5 + i].height = 30
tabelkop(ds, 8, 2, ["Code", "Eenheden per 24 u", "Zo nodig (J/N)", "Max. eenheden per 24 u", "Omschrijving"], 30)
for r in range(9, 59):
    for j in range(5):
        invoer(ds.cell(row=r, column=2 + j))
    ds.cell(row=r, column=6).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ds.row_dimensions[r].height = 20
ds["B9"], ds["C9"], ds["D9"], ds["F9"] = "VOORB1", 2, "N", "Voorbeeld (fictieve code, mag weg): 2 eenheden per 24 u, vast gebruik"
dv_jn = DataValidation(type="list", formula1='"J,N"', allow_blank=True)
ds.add_data_validation(dv_jn)
dv_jn.add("D9:D58")
ds.freeze_panes = "A9"
ds.sheet_properties.tabColor = GROEN
ds.page_setup.fitToWidth = 1
ds.page_setup.fitToHeight = 0
ds.sheet_properties.pageSetUpPr.fitToPage = True

# ================================================================== Toelichting
ts = wb.create_sheet("Toelichting")
ts.column_dimensions["A"].width = 2
ts.column_dimensions["B"].width = 118
ts.column_dimensions["C"].width = 2
kopbalk(ts, "Toelichting", "Werkwijze, rekenmethode, bronnen en verantwoording", "C")
ts.sheet_properties.tabColor = GRIJS
ts.page_setup.fitToWidth = 1
ts.page_setup.fitToHeight = 0
ts.sheet_properties.pageSetUpPr.fitToPage = True
H = "kop"
lines = [
    (H, "Werkwijze"),
    ("", "1.  Uitlezen: open in Pharmacom het dossier van de patiënt en druk op Ctrl+Alt+F3 (ook vanuit Excel). Pharmacom wordt "
         "dan zelf gekozen, zo nodig wordt de medicatiestatus (F3) geopend, en stap 1 wordt ingevuld. Handmatig invullen kan ook: "
         "kies per regel het middel en vul de totale dosis van de afgelopen 24 uur in (pleisters en continue fentanyl in µg/uur)."),
    ("", "2.  Kies de reden van rotatie (standaard: onvoldoende pijnstilling → 100%; bijwerkingen → 75%) en eventueel een extra reductie."),
    ("", "3.  Kies de concentratie morfine (10 of 1 mg/ml, eigen voorraad) en vul de kleinste instelstap van de pomp in. "
         "De berekening laat ook de pompstand bij de andere concentratie zien; bij een lage dosis (te grove afronding) of een "
         "spuit die binnen 24 uur leeg is, stelt de rekenhulp de andere concentratie voor."),
    ("", "4.  Lees de pompstand, dagdosis, doorbraakbolus en looptijd af in de tegels; de berekening staat eronder."),
    ("", "5.  Lees altijd de controles en waarschuwingen. Gebruik daarna 'Wis invoer' en sla het bestand niet op met patiëntgegevens."),
    ("", "Omzetten naar een ander opioïd (bijv. oraal of pleister in plaats van de pomp): tabblad Omzetten. Kies het nieuwe middel; "
         "de rekenhulp neemt de OME en het percentage uit stap 1 en 2 over, rekent de nieuwe dosis uit en kiest uit tabblad "
         "Preparaten de best passende sterkte (of combinatie van twee sterktes) met doseercode, plus doorbraakmedicatie "
         "(1/6 van de dagdosis). Het assortiment op tabblad Preparaten is aan te passen aan de eigen voorraad."),
    ("", ""),
    (H, "Uitlezen uit Pharmacom"),
    ("", "Op veel computers mag Excel zelf geen programma's starten (beveiliging tegen macrovirussen). Daarom start het uitlezen "
         "buiten Excel, via de sneltoets Ctrl+Alt+F3 of de snelkoppeling 'Lees opioiden uit Pharmacom' op het bureaublad. "
         "Het script leest via de Java Access Bridge (alleen lezen), herkent opioïden aan de ATC-code, de sterkte aan de "
         "etiketnaam en vertaalt de doseercode. De kolom 'Uitgelezen uit Pharmacom' laat per regel zien wat er gelezen en hoe het vertaald is."),
    ("", "Zo nodig / bij pijn: wordt als 'Doorbraak' gezet met het MAXIMUM per 24 u uit de doseercode (bijv. P1CM6 = max. 6 per dag). "
         "Verlaag dit naar het werkelijke gebruik als dat minder is; staat er geen maximum in de code, dan blijft de dosis leeg. "
         "Oranje = controleren. Niet automatisch: buprenorfine sublinguaal, onbekende codes."),
    ("", "Nierfunctie: na de medicatiestatus opent het script het Diagnostisch dossier (Alt+I, Alt+D), neemt de nieuwste eGFR "
         "uit Laboratorium uitslagen over in stap 2 (bij voorkeur CKD-EPI; datum in de opmerking bij de cel; oranje als ouder dan "
         "3 maanden of een grenswaarde zoals >90) en gaat terug naar de medicatiestatus (F3). Geen eGFR gevonden: het veld blijft leeg."),
    ("", "Bovenaan staat welk dossier is uitgelezen - controleer of dit de juiste patiënt is. Pharmacom_naar_Excel.ps1 en "
         "Pharmacom_medicatiestatus.ps1 moeten in dezelfde map staan als deze werkmap."),
    ("", ""),
    (H, "Rekenmethode"),
    ("", "OME (mg morfine oraal/24 u) = Σ dosis × factor  →  morfine s.c./i.v. = OME ÷ verhouding oraal : s.c. (standaard 3)  →  "
         "× % volgens reden × (1 − extra reductie)  →  ÷ 24 = mg/uur  →  ÷ concentratie = ml/uur."),
    ("", "Voorbeeld (ingevuld): oxycodon oraal 80 + 30 mg = 110 mg × 1,5 = 165 mg OME → ÷ 3 = 55 mg morfine s.c./24 u bij "
         "onvoldoende pijnstilling (100%) = 2,29 mg/uur → bij 10 mg/ml: 0,23 ml/uur. Bij bijwerkingen (75%): 41,3 mg/24 u = 0,17 ml/uur."),
    ("", ""),
    (H, "Kleuren"),
    ("", "Wit vlak met blauwe rand = invulveld of aanpasbare aanname.  Lichtgrijs = berekend.  Groen = belangrijkste uitkomst.  "
         "Oranje = controleren of waarschuwing."),
    ("", ""),
    (H, "Bronnen"),
    ("", "• Hoofdbron: KNMP Informatorium Medicamentorum, inleidende tekst Analgetica (G806), rubriek Dosering - Opioïdrotatie, "
         "conversiefactoren, opioïdrotatietabel en Doorbraakpijn (geraadpleegd 14-10-2026). Deze factoren zijn afkomstig uit de "
         "NVA-richtlijn Pijn bij patiënten met kanker; de NHG-Standaard Pijn hanteert dezelfde omrekening."),
    ("", "• KNMP: rotatie wegens bijwerkingen → start 75% van de equivalente 24-uursdosis; wegens onvoldoende effect → equivalente dosis. "
         "Doorbraak: 1/6 van de dagdosis, max. 4-6× per dag; > 3× per dag nodig → onderhoudsdosis met 50% verhogen."),
    ("", "• Aanvullend: NVA-richtlijn Pijn bij COPD/hartfalen, module Opioïdrotatie (2019) - max. 140 µg/uur buprenorfine, max. tapentadol; "
         "Pallialine-pijnkaart - doorbraak 10-15% (lagere variant)."),
    ("", "• Doseercodes: Centrale doseertabel, Opbouw doseercodes (Health Base, helpkaart Pharmacom/Medicom)."),
    ("", "• Codeïne staat niet in de KNMP-rotatietabel (literatuurwaarde 1 : 10)."),
    ("", "• Preparaten (tabblad Preparaten / Omzetten): KNMP Kennisbank, Voorschrijfproducten - producten uit de G-Standaard op "
         "PRK-niveau (geraadpleegd 15-10-2026). Doseerfrequentie MGA standaard 2x per dag; pas aan waar nodig."),
    ("", "• Methadon (dosisafhankelijk: 10 mg = 40 mg morfine oraal bij 0-20 mg/24 u, 60 mg bij 21-50 mg, 80 mg bij > 50 mg) en "
         "fentanyl oromucosaal (200 µg = 10 mg) en nasaal (100 µg = 30 mg): Equivalentietabel opioïden, Zaans Medisch Centrum "
         "(2020), met als bronnen Pallialine (methadon), Perioperative Medicine for the Junior Clinician (oromucosaal) en het "
         "EPAR van Instanyl (nasaal). KNMP noemt voor methadon alleen 'variabel 5 : 1 tot 10 : 1'."),
    ("", ""),
    (H, "Verantwoording"),
    ("", "Equianalgetische tabellen zijn benaderingen met beperkt wetenschappelijk bewijs; incomplete kruistolerantie, nier- en leverfunctie, "
         "leeftijd en comedicatie beïnvloeden de juiste dosis. Deze rekenhulp ondersteunt maar vervangt geen klinisch oordeel, lokaal "
         "protocol of dubbele controle. Bij twijfel: consultatieteam palliatieve zorg of pijnteam."),
    ("", ""),
    ("klein", VERSIE + "  ·  factoren volgens KNMP IM  ·  uitlezen uit Pharmacom  ·  github.com/ronaldsteenkamp/opioidrotatie"),
]
for i, (soort, t) in enumerate(lines):
    r = 5 + i
    c = ts.cell(row=r, column=2, value=t)
    if soort == H:
        c.font = Font(name=FS, size=12, color=BLAUW)
        c.border = Border(bottom=side(GROEN, "medium"))
        ts.row_dimensions[r].height = 24
    elif soort == "klein":
        c.font = font(8, BLAUW_M, italic=True)
    else:
        c.font = font(10, GRIJS)
        regels = max(1, len(t) // 130 + 1)
        ts.row_dimensions[r].height = 16 * regels + 4 if t else 8
    c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)

# ================================================================== Preparaten
# Assortiment per middel: welke sterktes er zijn en hoe ze gedoseerd worden. Tabblad Omzetten
# kiest hieruit het preparaat en de doseercode. Aan te passen aan het eigen assortiment.
ps = wb.create_sheet("Preparaten")
for col, w in zip("ABCDEFGHIJK", [2, 22, 34, 10, 10, 12, 10, 14, 18, 44, 2]):
    ps.column_dimensions[col].width = w
kopbalk(ps, "Preparaten", "Sterktes en doseercodes voor het advies op tabblad Omzetten. Producten volgens KNMP Kennisbank, "
        "Voorschrijfproducten (G-Standaard, PRK-niveau; 15-10-2026). Pas aan aan het eigen assortiment.", "K")
ps["B5"] = ("Doseercode: {f} = aantal keer per dag, {n} = aantal eenheden per keer (Centrale doseertabel, bijv. 2.1T = 2x per dag "
            "1 tablet; D31PL = per 3 dagen 1 pleister; P1CM6 = bij pijn 1 capsule, max. 6 per dag; {m} = 6 x {n}, dus max. 6 keer per dag). Keer per 24 u leeg = pleister: "
            "de dosis is dan de sterkte in µg/uur. 'Rekenen als' (doorbraak): middel waarvan de omrekenfactor gebruikt wordt.")
ps["B5"].font = font(9, GRIJS)
ps["B5"].alignment = boven
ps.merge_cells("B5:J5")
ps.row_dimensions[5].height = 44
tabelkop(ps, 6, 2, ["Middel (omzetten naar)", "Preparaat", "Sterkte", "Eenheid", "Gebruik", "Keer per 24 u",
                    "Doseercode", "Rekenen als", "Opmerking"], 30)
O, B_ = "Onderhoud", "Doorbraak"
prep = []


def nl(s):
    return str(s).replace(".", ",")


# Productnamen volgens KNMP Kennisbank, Voorschrijfproducten (G-Standaard, PRK-niveau), geraadpleegd 15-10-2026
for s in (10, 15, 30, 60, 100):
    prep.append(("Morfine oraal", f"MORFINE TABLET MGA {s}MG (SULFAAT)", s, "mg", O, 2, "{f}.{n}T", "", ""))
for s in (5, 10, 15, 20, 30, 40, 60, 80, 120):
    prep.append(("Oxycodon oraal", f"OXYCODON TABLET MGA {s}MG", s, "mg", O, 2, "{f}.{n}T", "", ""))
for s in (4, 8, 16, 24):
    prep.append(("Hydromorfon oraal", f"HYDROMORFON CAPSULE MGA {s}MG", s, "mg", O, 2, "{f}.{n}C", "", ""))
for s in (12, 25, 37, 50, 75, 100):
    prep.append(("Fentanyl pleister", f"FENTANYL PLEISTER {s} UG/UUR", s, "µg/uur", O, None, "D3{n}PL", "",
                 "Elke 72 uur vervangen"))
for s in (5, 10, 15, 20):
    prep.append(("Buprenorfine pleister", f"BUPRENORFINE PLEISTER {s}UG/UUR 7 DAGEN", s, "µg/uur", O, None, "1W{n}PL", "",
                 "Elke 7 dagen vervangen"))
for s in (35, 52.5, 70):
    prep.append(("Buprenorfine pleister", f"BUPRENORFINE PLEISTER {nl(s)}UG/UUR 3 DAGEN", s, "µg/uur", O, None, "D3{n}PL", "",
                 "Elke 3 dagen vervangen"))
for s in (35, 52.5, 70):
    prep.append(("Buprenorfine pleister", f"BUPRENORFINE PLEISTER {nl(s)}UG/UUR 4 DAGEN", s, "µg/uur", O, None,
                 "D4{n}PL", "", "Elke 4 dagen vervangen"))
for s in (100, 150, 200):
    prep.append(("Tramadol oraal", f"TRAMADOL TABLET MGA {s}MG", s, "mg", O, 2, "{f}.{n}T", "", ""))
for s in (50, 100, 150):
    prep.append(("Tramadol oraal", f"TRAMADOL CAPSULE MGA {s}MG", s, "mg", O, 2, "{f}.{n}C", "", ""))
for s in (25, 50, 100, 150, 200, 250):
    prep.append(("Tapentadol oraal", f"TAPENTADOL TABLET MGA {s}MG", s, "mg", O, 2, "{f}.{n}T", "", ""))
for s in (2, 5, 10, 20, 40, 50):
    prep.append(("Methadon oraal", f"METHADON TABLET {s}MG", s, "mg", O, 2, "{f}.{n}T", "", "Alleen in overleg met deskundige"))
# Doorbraak: snelwerkend preparaat; bij pleisters snelwerkende morfine. Tapentadol heeft geen snelwerkende vorm.
for doel in ("Morfine oraal", "Fentanyl pleister", "Buprenorfine pleister"):
    for naam, s in (("MORFINE TABLET 10MG (SULFAAT)", 10), ("MORFINE TABLET 20MG", 20)):
        prep.append((doel, naam, s, "mg", B_, None, "P{n}TM{m}", "Morfine oraal", ""))
for s in (5, 10, 20):
    prep.append(("Oxycodon oraal", f"OXYCODON CAPSULE {s}MG", s, "mg", B_, None, "P{n}CM{m}", "Oxycodon oraal", ""))
for s in (1.3, 2.6):
    prep.append(("Hydromorfon oraal", f"HYDROMORFON CAPSULE {nl(s)}MG", s, "mg", B_, None, "P{n}CM{m}", "Hydromorfon oraal", ""))
prep.append(("Tramadol oraal", "TRAMADOL CAPSULE 50MG", 50, "mg", B_, None, "P{n}CM{m}", "Tramadol oraal", ""))
prep.sort(key=lambda p: ([d[0] for d in drugs].index(p[0]), p[4] != O))
P0, P1 = 7, 86
dv_doel = DataValidation(type="list", formula1=NAMES, allow_blank=True)
dv_gebruik = DataValidation(type="list", formula1=f'"{O},{B_}"', allow_blank=True)
dv_als = DataValidation(type="list", formula1=NAMES, allow_blank=True)
for v in (dv_doel, dv_gebruik, dv_als):
    ps.add_data_validation(v)
for i, r in enumerate(range(P0, P1 + 1)):
    rij = prep[i] if i < len(prep) else (None,) * 9
    for j, v in enumerate(rij):
        c = ps.cell(row=r, column=2 + j)
        invoer(c, v)
        c.alignment = links if j in (0, 1, 6, 7, 8) else Alignment(horizontal="center", vertical="center")
    ps.cell(row=r, column=4).number_format = "General"
    ps.row_dimensions[r].height = 20
dv_doel.add(f"B{P0}:B{P1}")
dv_gebruik.add(f"F{P0}:F{P1}")
dv_als.add(f"I{P0}:I{P1}")
ps.freeze_panes = f"A{P0}"
ps.sheet_properties.tabColor = GROEN
ps.page_setup.orientation = "landscape"
ps.page_setup.fitToWidth = 1
ps.page_setup.fitToHeight = 0
ps.sheet_properties.pageSetUpPr.fitToPage = True


def prep_kol(k):
    return f"Preparaten!${k}${P0}:${k}${P1}"


PB, PC, PD, PF, PG, PH, PI = (prep_kol(k) for k in "BCDFGHI")


def getal(x):
    """Excel-formule (zonder '='): getal als tekst met 0-2 decimalen, landinstelling-onafhankelijk."""
    return f"IF(ROUND({x},0)={x},FIXED({x},0,TRUE),IF(ROUND({x},1)={x},FIXED({x},1,TRUE),FIXED({x},2,TRUE)))"


def procent(x):
    return f'IF({x}>=0,"+","")&FIXED({x}*100,0,TRUE)&"%"'


# ================================================================== Omzetten
zs = wb.create_sheet("Omzetten")
for col, w in zip("ABCDEFGHIJ", [2, 40, 13, 12, 11, 14, 14, 12, 40, 2]):
    zs.column_dimensions[col].width = w
for col in "KLMNOPQ":
    zs.column_dimensions[col].hidden = True
kopbalk(zs, "Omzetten naar ander opioïd", "Equivalente dosis, preparaat en doseercode op basis van de huidige medicatie "
        "(tabblad Rotatie, stap 1 en 2) en het assortiment op tabblad Preparaten.", "J")
zs.sheet_view.showRowColHeaders = False
zs.sheet_properties.tabColor = GROEN
zs.merge_cells("B4:I4")
zs["B4"] = "=Rotatie!B4"
zs["B4"].font = Font(name=FS, size=10, color=GROEN_D)
zs["B4"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
for c in "BCDEFGHI":
    zs[f"{c}4"].fill = fill(GROEN_L)
zs.row_dimensions[4].height = 26


def zkop(r, nummer, titel):
    c = zs[f"B{r}"]
    c.value = f"{nummer}   {titel}"
    c.font = Font(name=FS, size=13, color=BLAUW)
    c.alignment = Alignment(vertical="center", indent=1)
    for k in "BCDEFGHI":
        zs[f"{k}{r}"].border = Border(bottom=side(GROEN, "medium"))
    zs.row_dimensions[r].height = 28


def zlabel(r, tekst, noot=None):
    zs[f"B{r}"] = tekst
    zs[f"B{r}"].font = f_label
    zs[f"B{r}"].alignment = links
    if noot:
        zs[f"E{r}"] = noot
        zs[f"E{r}"].font = f_klein
        zs[f"E{r}"].alignment = links
        zs.merge_cells(f"E{r}:I{r}")
    zs.row_dimensions[r].height = 22


def zwaarde(ref, formule, fmt="@", sterk=False):
    c = zs[ref]
    c.value = formule
    c.number_format = fmt
    c.font = f_sterk if sterk else f_waarde
    c.fill = fill(LICHT)
    c.alignment = rechts


# 1. huidige medicatie (uit tabblad Rotatie)
zkop(6, 1, "Huidige medicatie (tabblad Rotatie, stap 1)")
tabelkop(zs, 7, 2, ["Middel + toedieningsweg", "Dosis per 24 u", "Eenheid", "Soort", "OME (mg/24 u)"], 26)
for i in range(R1 - R0 + 1):
    r, rr = 8 + i, R0 + i
    zs[f"B{r}"] = f'=IF(Rotatie!C{rr}="","",Rotatie!C{rr})'
    zs[f"C{r}"] = f'=IF(Rotatie!C{rr}="","",Rotatie!D{rr})'
    zs[f"D{r}"] = f'=IF(Rotatie!C{rr}="","",Rotatie!E{rr})'
    zs[f"E{r}"] = f'=IF(Rotatie!C{rr}="","",Rotatie!B{rr})'
    zs[f"F{r}"] = f'=IF(Rotatie!C{rr}="","",Rotatie!G{rr})'
    for k in "BCDEF":
        zs[f"{k}{r}"].font = f_norm
        zs[f"{k}{r}"].border = rand_cel
        zs[f"{k}{r}"].alignment = links if k == "B" else midden
    zs[f"C{r}"].number_format = zs[f"F{r}"].number_format = "0.0##"
    zs.row_dimensions[r].height = 20
ZT = 8 + R1 - R0 + 1                        # 14
zlabel(ZT, "Totaal OME (mg morfine oraal / 24 u)")
zwaarde(f"C{ZT}", f"=Rotatie!G{TR}", "0.0", sterk=True)
zlabel(ZT + 1, "Toe te passen %", "x")
zs[f"E{ZT + 1}"] = '="Reden: "&Rotatie!D19&IF(N(Rotatie!D21)>0,", extra reductie "&FIXED(Rotatie!D21*100,0,TRUE)&"%","")&"  (tabblad Rotatie, stap 2)"'
zwaarde(f"C{ZT + 1}", "=Rotatie!D23", "0%")
zlabel(ZT + 2, "Toegepaste OME (mg/24 u)")
zwaarde(f"C{ZT + 2}", f"=C{ZT}*C{ZT + 1}", "0.0", sterk=True)
OME, OMET = f"$C${ZT}", f"$C${ZT + 2}"

# 2. keuze
Z2 = ZT + 4                                 # 18
zkop(Z2, 2, "Omzetten naar")
zlabel(Z2 + 1, "Nieuw middel + toedieningsweg", "Keuzelijst (tabblad Omrekenfactoren).")
invoer(zs[f"C{Z2 + 1}"], "Fentanyl pleister")
zs.merge_cells(f"C{Z2 + 1}:D{Z2 + 1}")
dv_z = DataValidation(type="list", formula1=NAMES, allow_blank=True,
                      error="Kies een middel uit de lijst (tabblad Omrekenfactoren).", showErrorMessage=True)
zs.add_data_validation(dv_z)
dv_z.add(f"C{Z2 + 1}")
zlabel(Z2 + 2, "Aantal keer per dag (optioneel)", "Leeg = standaard uit tabblad Preparaten (bijv. MGA 2x per dag).")
invoer(zs[f"C{Z2 + 2}"], None, "0")
dv_f = DataValidation(type="whole", operator="between", formula1="1", formula2="6", allow_blank=True,
                      error="Voer 1 tot 6 keer per dag in.", showErrorMessage=True)
zs.add_data_validation(dv_f)
dv_f.add(f"C{Z2 + 2}")
DOEL, FREQ = f"$C${Z2 + 1}", f"$C${Z2 + 2}"
# Verborgen volgnummers op Preparaten: de k-de regel voor het gekozen middel (zonder matrixformules)
for kol, gebruik in (("L", O), ("M", B_)):
    ps.column_dimensions[kol].hidden = True
    for r in range(P0, P1 + 1):
        ps[f"{kol}{r}"] = (f'=IF(AND($B{r}=Omzetten!{DOEL},$F{r}="{gebruik}"),'
                           f'COUNTIFS($B${P0}:$B{r},Omzetten!{DOEL},$F${P0}:$F{r},"{gebruik}"),"")')
FT = f"INDEX({FACT},MATCH({DOEL},{NAMES},0))"
UT = f"INDEX({UNITS},MATCH({DOEL},{NAMES},0))"
MAXT = f"INDEX({MAXD},MATCH({DOEL},{NAMES},0))"


def dagdosis(ome):
    return (f'=IF(OR({DOEL}="",N({ome})=0),"",IFERROR(IF({DOEL}="Methadon oraal",{methadon_dosis(ome)},'
            f'{ome}/{FT}),""))')


zlabel(Z2 + 3, "Equivalente dosis (100%)")
zwaarde(f"C{Z2 + 3}", dagdosis(OME), "0.0#")
zlabel(Z2 + 4, "Nieuwe dosis (toegepast %)")
zwaarde(f"C{Z2 + 4}", dagdosis(OMET), "0.0#", sterk=True)
for r in (Z2 + 3, Z2 + 4):
    zs[f"D{r}"] = f'=IF({DOEL}="","",{UT})'
    zs[f"D{r}"].font = f_klein
    zs[f"D{r}"].alignment = links
zs[f"E{Z2 + 4}"] = (f'=IF(AND(ISNUMBER(C{Z2 + 4}),{DOEL}<>""),IF(RIGHT({UT},4)="/uur","per uur (pleister/continu)",'
                    f'"per 24 uur = "&{getal(f"C{Z2 + 4}/24")}&" "&LEFT({UT},2)&"/uur"),"")')
zs[f"E{Z2 + 4}"].font = f_klein
zs.merge_cells(f"E{Z2 + 4}:I{Z2 + 4}")
DAG = f"$C${Z2 + 4}"


def lijst(r0, n, gebruik):
    """Regels van tabblad Preparaten voor het gekozen middel; verborgen kolom K = regelnummer."""
    volg = prep_kol("L" if gebruik == O else "M")
    for k in range(1, n + 1):
        r = r0 + k - 1
        zs[f"K{r}"] = f'=IFERROR(MATCH({k},{volg},0),"")'
        zs[f"B{r}"] = f'=IF(K{r}="","",INDEX({PC},K{r}))'
        zs[f"C{r}"] = f'=IF(K{r}="","",INDEX({PD},K{r}))'
        zs[f"N{r}"] = f'=IF(K{r}="","",INDEX({PH},K{r}))'
        for kol in "BCDEFGH":
            zs[f"{kol}{r}"].font = f_norm
            zs[f"{kol}{r}"].border = rand_cel
            zs[f"{kol}{r}"].alignment = links if kol == "B" else midden
        zs[f"C{r}"].number_format = "General"
        zs[f"H{r}"].number_format = "+0%;-0%;0%"
        zs.row_dimensions[r].height = 20


def code(r, n):
    return f'SUBSTITUTE(SUBSTITUTE(N{r},"{{f}}",D{r}),"{{n}}",{n})'


# 3. onderhoud
Z3 = Z2 + 6                                 # 24
zkop(Z3, 3, "Advies onderhoud")
zlabel(Z3 + 1, "Eén sterkte")
zlabel(Z3 + 2, "Combinatie van twee sterktes")
L0, LN = Z3 + 4, 12                         # lijst 28..39
tabelkop(zs, Z3 + 3, 2, ["Preparaat", "Sterkte", "Keer per 24 u", "Per keer", "Doseercode", "Dosis", "Afwijking"], 26)
lijst(L0, LN, O)
for k in range(LN):
    r = L0 + k
    zs[f"D{r}"] = f'=IF(K{r}="","",IF(INDEX({PG},K{r})="",1,IF(N({FREQ})>0,{FREQ},INDEX({PG},K{r}))))'
    zs[f"L{r}"] = f'=IF(OR(K{r}="",NOT(ISNUMBER({DAG}))),"",{DAG}/D{r})'          # doel per keer
    zs[f"E{r}"] = f'=IF(L{r}="","",ROUND(L{r}/C{r},0))'
    zs[f"F{r}"] = f'=IF(E{r}="","",IF(E{r}<1,"te laag",{code(r, f"E{r}")}))'
    zs[f"G{r}"] = f'=IF(E{r}="","",IF(E{r}<1,"",E{r}*C{r}*D{r}))'
    zs[f"H{r}"] = f'=IF(N(G{r})=0,"",G{r}/{DAG}-1)'
    zs[f"M{r}"] = f'=IF(E{r}="","",IF(OR(E{r}<1,E{r}>4),"",ABS(H{r})+0.03*(E{r}-1)))'   # score: afwijking, weinig eenheden
    zs[f"G{r}"].number_format = "0.0#"
rng = lambda kol: f"${kol}${L0}:${kol}${L0 + LN - 1}"   # noqa: E731
ZB = Z3 + 1                                 # beste regel (één sterkte)
zs[f"K{ZB}"] = f'=IFERROR(MATCH(MIN({rng("M")}),{rng("M")},0),"")'
zs[f"L{ZB}"] = f'=IF(K{ZB}="","",INDEX({rng("H")},K{ZB}))'
zs[f"C{ZB}"] = (f'=IF({DOEL}="","",IF(N({OME})=0,"Vul eerst stap 1 in op tabblad Rotatie.",'
                f'IF(COUNT({rng("C")})=0,"Geen preparaten voor dit middel op tabblad Preparaten.",'
                f'IF(K{ZB}="","Dosis past niet bij de beschikbare sterktes - zie de tabel hieronder.",'
                f'INDEX({rng("B")},K{ZB})&":  "&INDEX({rng("F")},K{ZB})&"   =  "&{getal(f"INDEX({rng("G")},K{ZB})")}&" "&{UT}'
                f'&"  ("&{procent(f"L{ZB}")}&")"))))')
# Combinatie: grootste sterkte <= dosis per keer (zelfde doseercode als de beste), aangevuld met de dichtstbijzijnde
ZC = Z3 + 2
P_, N_ = f"INDEX({rng('L')},$K${ZB})", f"INDEX({rng('N')},$K${ZB})"
GROOTSTE = f'_xlfn.MAXIFS({rng("C")},{rng("C")},"<="&{P_}*1.0001,{rng("N")},{N_})'
zs[f"K{ZC}"] = f'=IF($K${ZB}="","",IF({GROOTSTE}=0,"",{GROOTSTE}))'
zs[f"L{ZC}"] = f'=IF(K{ZC}="","",INT({P_}/K{ZC}+0.000001))'                           # aantal grote
zs[f"M{ZC}"] = f'=IF(K{ZC}="","",{P_}-L{ZC}*K{ZC})'                                    # rest per keer
zs[f"N{ZC}"] = f'=IF(K{ZC}="","",MIN({rng("O")}))'
for k in range(LN):   # afstand van elke sterkte (zelfde doseercode) tot de rest
    r = L0 + k
    zs[f"O{r}"] = f'=IF(OR($K${ZC}="",K{r}=""),"",IF(N{r}<>{N_},"",ABS(C{r}-$M${ZC})))'
zs[f"O{ZC}"] = f'=IF(K{ZC}="","",IF(M{ZC}<=N{ZC},0,_xlfn.MINIFS({rng("C")},{rng("O")},N{ZC})))'   # aanvulling
COMBI = f'(L{ZC}*K{ZC}+O{ZC})*INDEX({rng("D")},$K${ZB})'
zs[f"Q{ZB}"] = f'=IF(K{ZB}="","",INDEX({rng("E")},K{ZB}))'                     # eenheden per keer (één sterkte)
zs[f"Q{ZC}"] = f'=IF(K{ZC}="","",L{ZC}+IF(N(O{ZC})>0,1,0))'                     # eenheden per keer (combinatie)
# Alleen tonen als de combinatie iets oplevert: andere dosis, of dezelfde dosis met minder eenheden
zs[f"P{ZC}"] = (f'=IF(OR(K{ZC}="",N(O{ZC})=0,O{ZC}=K{ZC}),"",IF(AND(ABS({COMBI}-INDEX({rng("G")},$K${ZB}))<0.0001,'
                f'Q{ZC}>=Q{ZB}),"",{COMBI}/{DAG}-1))')
GROOT = f'INDEX({rng("B")},MATCH(K{ZC},{rng("C")},0))'
KLEIN = f'INDEX({rng("B")},MATCH(O{ZC},{rng("C")},0))'
CODE_G = (f'SUBSTITUTE(SUBSTITUTE({N_},"{{f}}",INDEX({rng("D")},$K${ZB})),"{{n}}",L{ZC})')
CODE_K = (f'SUBSTITUTE(SUBSTITUTE({N_},"{{f}}",INDEX({rng("D")},$K${ZB})),"{{n}}",1)')
zs[f"C{ZC}"] = (f'=IF(OR({DOEL}="",$K${ZB}=""),"",IF(P{ZC}="","Niet nodig - één sterkte past.",'
                f'{GROOT}&":  "&{CODE_G}&"   +   "&{KLEIN}&":  "&{CODE_K}&"   =  "'
                f'&{getal(f"(L{ZC}*K{ZC}+O{ZC})*INDEX({rng('D')},$K${ZB})")}&" "&{UT}&"  ("&{procent(f"P{ZC}")}&")"))')
for r in (ZB, ZC):
    zs.merge_cells(f"C{r}:I{r}")
    zs[f"C{r}"].font = f_sterk
    zs[f"C{r}"].fill = fill(LICHT)
    zs[f"C{r}"].alignment = links
    zs.row_dimensions[r].height = 26
# Het advies met de kleinste afwijking groen
COMBI_BETER = (f'IF(ISNUMBER($P${ZC}),IF(ISNUMBER($L${ZB}),OR(ABS($P${ZC})<ABS($L${ZB})-0.001,'
               f'AND(ABS($P${ZC})<=ABS($L${ZB})+0.001,$Q${ZC}<$Q${ZB})),TRUE),FALSE)')
zs.conditional_formatting.add(f"C{ZB}", FormulaRule(
    formula=[f'AND(ISNUMBER($L${ZB}),NOT({COMBI_BETER}))'], font=Font(color=GROEN_D, bold=True), fill=cf_fill(GROEN_L)))
zs.conditional_formatting.add(f"C{ZC}", FormulaRule(
    formula=[COMBI_BETER], font=Font(color=GROEN_D, bold=True), fill=cf_fill(GROEN_L)))

# 4. doorbraak
Z4 = L0 + LN + 1                            # 41
zkop(Z4, 4, "Doorbraakmedicatie")
zlabel(Z4 + 1, "Doorbraakdosis per keer (als OME)", "Fractie van de dagdosis: tabblad Omrekenfactoren (KNMP 1/6), max. 4-6x per dag.")
zwaarde(f"C{Z4 + 1}", f'=IF(N({OMET})=0,"",{OMET}*{S_BOLLO})', "0.0")
zs[f"D{Z4 + 1}"] = "mg"
zs[f"D{Z4 + 1}"].font = f_klein
zlabel(Z4 + 2, "Advies doorbraak")
B0, BN = Z4 + 4, 4                          # lijst 45..48
tabelkop(zs, Z4 + 3, 2, ["Preparaat", "Sterkte", "Rekenen als", "Per keer", "Doseercode", "Dosis per keer", "Afwijking"], 26)
lijst(B0, BN, B_)
BTOME = f"$C${Z4 + 1}"
for k in range(BN):
    r = B0 + k
    zs[f"D{r}"] = f'=IF(K{r}="","",INDEX({PI},K{r}))'
    zs[f"L{r}"] = f'=IF(OR(K{r}="",N({BTOME})=0),"",IFERROR({BTOME}/INDEX({FACT},MATCH(D{r},{NAMES},0)),""))'
    zs[f"E{r}"] = f'=IF(L{r}="","",ROUND(L{r}/C{r},0))'
    zs[f"F{r}"] = f'=IF(E{r}="","",IF(E{r}<1,"te laag",SUBSTITUTE(SUBSTITUTE(N{r},"{{n}}",E{r}),"{{m}}",6*E{r})))'
    zs[f"G{r}"] = f'=IF(E{r}="","",IF(E{r}<1,"",E{r}*C{r}))'
    zs[f"H{r}"] = f'=IF(N(G{r})=0,"",G{r}/L{r}-1)'
    zs[f"M{r}"] = f'=IF(E{r}="","",IF(OR(E{r}<1,E{r}>2),"",ABS(H{r})+0.03*(E{r}-1)))'
    zs[f"D{r}"].font = f_klein
    zs[f"G{r}"].number_format = "0.0#"
brng = lambda kol: f"${kol}${B0}:${kol}${B0 + BN - 1}"   # noqa: E731
ZD = Z4 + 2
zs[f"K{ZD}"] = f'=IFERROR(MATCH(MIN({brng("M")}),{brng("M")},0),"")'
zs[f"C{ZD}"] = (f'=IF(OR({DOEL}="",N({OMET})=0),"",IF(COUNT({brng("C")})=0,"Geen doorbraakpreparaat voor dit middel op tabblad Preparaten.",'
                f'IF(K{ZD}="","Doorbraakdosis lager dan de kleinste sterkte - zie de tabel hieronder.",'
                f'INDEX({brng("B")},K{ZD})&":  "&INDEX({brng("F")},K{ZD})&"   =  "&{getal(f"INDEX({brng("G")},K{ZD})")}'
                f'&" mg per keer  ("&{procent(f"INDEX({brng('H')},K{ZD})")}&")")))')
zs.merge_cells(f"C{ZD}:I{ZD}")
zs[f"C{ZD}"].font = Font(name=FS, size=10, color=GROEN_D)
zs[f"C{ZD}"].fill = fill(GROEN_L)
zs[f"C{ZD}"].alignment = links
zs.row_dimensions[ZD].height = 26

# 5. controles
Z5 = B0 + BN + 1                            # 50
zkop(Z5, 5, "Controles en waarschuwingen")
zwarns = [
    f'=IF(AND({DOEL}<>"",COUNTIF(Rotatie!C{R0}:C{R1},{DOEL})>0),"Het gekozen middel staat al in de huidige medicatie. '
    f'Bedoel je een ander middel?","")',
    f'=IF(AND(ISNUMBER({DAG}),ISNUMBER({MAXT})),IF({DAG}>{MAXT},"Nieuwe dosis boven de max. dagdosis ("&{MAXT}&" "&{UT}&").",""),"")',
    f'=IF(NOT(ISNUMBER($L${ZB})),"",IF(ABS($L${ZB})<=0.1,"",IF(ISNUMBER($P${ZC}),IF(ABS($P${ZC})<=0.1,"",'
    f'"Het advies wijkt meer dan 10% af van de berekende dosis - kies zelf een passende dosering."),'
    f'"Het advies wijkt meer dan 10% af van de berekende dosis - kies zelf een passende dosering.")))',
    f'=IF({DOEL}="Fentanyl pleister","Fentanylpleister: de werking begint pas na 12-24 uur. Bouw de vorige medicatie af '
    f'volgens protocol, zorg voor doorbraakmedicatie en pas de dosis niet eerder aan dan na 3 dagen.","")',
    f'=IF({DOEL}="Buprenorfine pleister","Buprenorfinepleister: partiële agonist, werking na 12-24 uur; boven 140 µg/uur '
    f'niet geadviseerd. Plan de overgang in overleg.","")',
    f'=IF({DOEL}="Methadon oraal","Methadon: alleen in overleg met een deskundige (niet-lineaire omrekening, lange en '
    f'variabele halfwaardetijd, stapeling).","")',
    f'=IF(OR({DOEL}="Fentanyl oromucosaal",{DOEL}="Fentanyl nasaal"),"Snelwerkend fentanyl is doorbraakmedicatie en wordt '
    f'los getitreerd (start met de laagste sterkte), niet op basis van equivalentie.","")',
    f'=IF(ISNUMBER(SEARCH("s.c./i.v.",{DOEL})),"Parenteraal: voor een morfinepomp zie tabblad Rotatie; anders de dosis per 24 uur '
    f'verdelen of continu toedienen volgens protocol.","")',
    f'=IF({OME}>={S_HIGH},"Hoge dosis (OME ≥ "&{S_HIGH}&" mg/24 u): de omrekening is onzekerder - overweeg extra reductie '
    f'en overleg.","")',
    '="Rekenhulp - vervangt geen klinisch oordeel, lokaal protocol of dubbele controle. Controleer het advies altijd."',
]
# Elke controle in een verborgen cel (kolom K); zichtbaar alleen de van toepassing zijnde, onder elkaar
for i, f_ in enumerate(zwarns[:-1]):
    zs[f"K{Z5 + 1 + i}"] = f_
c = zs[f"B{Z5 + 1}"]
WLIJST = f'_xlfn.TEXTJOIN(CHAR(10),TRUE,K{Z5 + 1}:K{Z5 + len(zwarns) - 1})'
c.value = f'=IF({DOEL}="","",IF({WLIJST}="","Geen bijzonderheden.",{WLIJST}))'
c.font = Font(name=F, size=10, color=ORANJE)
c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
zs.merge_cells(f"B{Z5 + 1}:I{Z5 + 1}")
zs.row_dimensions[Z5 + 1].height = 96
ZEIND = Z5 + 2
zs[f"B{ZEIND}"] = zwarns[-1]
zs[f"B{ZEIND}"].font = font(9, GRIJS, italic=True)
zs.merge_cells(f"B{ZEIND}:I{ZEIND}")
zs.print_area = f"A1:J{ZEIND}"
zs.page_setup.orientation = "portrait"
zs.page_setup.paperSize = zs.PAPERSIZE_A4
zs.page_setup.fitToWidth = 1
zs.page_setup.fitToHeight = 1
zs.sheet_properties.pageSetUpPr.fitToPage = True
zs.freeze_panes = "A5"

# Volgorde tabbladen: Rotatie, Omzetten, Omrekenfactoren, Preparaten, Doseercodes, Omrekentabel, Toelichting
wb._sheets = [ws, zs, fs, ps, ds, qs, ts]
wb.active = 0

# Samenvoegen wist de opmaak van de cellen achter de eerste cel; zet de vulling terug,
# anders ontstaan er naden in gekleurde vlakken (tegels, patiëntregel).
for sh in wb.worksheets:
    for mr in sh.merged_cells.ranges:
        anker = sh.cell(mr.min_row, mr.min_col)
        if anker.fill is not None and anker.fill.fill_type:
            for rij in sh.iter_rows(min_row=mr.min_row, max_row=mr.max_row, min_col=mr.min_col, max_col=mr.max_col):
                for c in rij:
                    c.fill = copy(anker.fill)

# Overal Segoe UI (ook cellen zonder eigen opmaak)
for sh in wb.worksheets:
    for row in sh.iter_rows():
        for c in row:
            if c.value is not None and c.font is not None and c.font.name not in (F, FS, "Segoe UI Symbol"):
                c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, italic=c.font.italic,
                              color=c.font.color)

wb.save(OUT)
print("saved", OUT)
